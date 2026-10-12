#!/usr/bin/env python3
"""Build an editable 48 x 36 inch, three-column scientific poster outline.

Requires python-pptx, Pillow and Matplotlib. Figures are deliberately empty
placeholders; the stored performance values describe completed experiments.
"""
import argparse
import hashlib
import json
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
WIDTH, HEIGHT = 48, 36
NAVY, INK, GRAY = '17334B', '233645', '536879'
ACCENTS = ('2376A2', '147E7A', 'B55C30')
FONT_PATH = Path('/usr/share/fonts/dejavu/DejaVuSans.ttf')
if not FONT_PATH.exists():
    FONT_PATH = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def layout():
    shapes = []

    def rectangle(name, x, y, w, h, fill, line=None, dashed=False):
        shapes.append(dict(kind='rect', name=name, x=x, y=y, w=w, h=h,
                           fill=fill, line=line, dashed=dashed))

    def text(name, x, y, w, h, value, size=28, color=INK, bold=False, align='left'):
        shapes.append(dict(kind='text', name=name, x=x, y=y, w=w, h=h,
                           text=value, font=size, color=color, bold=bold, align=align))

    rectangle('title-banner', 0, 0, WIDTH, 3.8, NAVY)
    text('meeting-label', .95, .3, 45.5, .4, 'BMES 2026  |  EDITABLE POSTER OUTLINE', 20, 'B9D4E2', True)
    text('poster-title', .95, .9, 46, 2.05,
         'syncoCT: Controlled Digital Sibling Populations\nfor Coarctation-Focused AI Evaluation', 66, 'FFFFFF', True)
    text('author-placeholder', .95, 3.10, 46, .5,
         '[Authors]  |  [Affiliations]  |  [Contact]', 28, 'FFFFFF')

    cols = (.95, 16.58, 32.21)
    cw = 14.84
    heads = ('1  INTRODUCTION + METHODS', '2  VIRTUAL POPULATION', '3  AI PERFORMANCE + VALIDATION')
    for i, x in enumerate(cols):
        rectangle(f'column-header-{i+1}', x, 4.25, cw, .93, ACCENTS[i])
        text(f'column-heading-{i+1}', x+.2, 4.42, cw-.4, .65, heads[i], 35, 'FFFFFF', True)
    for x in (16.15, 31.78):
        rectangle(f'column-divider-{x}', x, 5.6, .015, 26.5, 'D6E1E6')

    x = cols[0]
    text('intro-heading', x, 5.6, cw, .55, 'THE DATA CHALLENGE', 31, ACCENTS[0], True)
    text('intro-copy', x, 6.32, cw, 2.15,
         'Robust AI needs controllable, consistently labeled data with deliberate coverage of uncommon anatomical configurations. Clinical cohort assembly can be limited by labeling effort and uneven disease coverage.', 29)
    text('hypothesis-heading', x, 8.78, cw, .55, 'RESEARCH HYPOTHESIS', 31, ACCENTS[0], True)
    text('hypothesis-copy', x, 9.48, cw, 2.08,
         'Anatomically diverse digital siblings may improve robustness on unseen anatomy and, ultimately, clinical CT. syncoCT makes this hypothesis testable through reproducible anatomical editing.', 28)

    x = cols[1]
    text('population-heading', x, 5.6, cw, .55, 'A CONTROLLED CoA-LIKE CASE STUDY', 31, ACCENTS[1], True)
    text('population-copy', x, 6.32, cw, 2.45,
         'Coarctation is congenital aortic narrowing; CDC estimates about 1 in 1,724 US births [2]. Generate isthmus/descending narrowing with companion arch narrowing and ascending/arch enlargement.', 29)
    rectangle('population-count-background', x, 9.4, cw, 1.9, 'E9F4F2')
    text('population-counts', x+.3, 9.68, cw-.6, 1.4,
         '12 whole-body XCAT phantom IDs [1]\n192 scenarios: 12 baselines + 180 edited siblings', 29, ACCENTS[1], True)

    x = cols[2]
    text('ai-heading', x, 5.6, cw, .55, 'IMAGE-BASED, COARSE-TO-FINE AI', 31, ACCENTS[2], True)
    text('ai-copy', x, 6.32, cw, 2.1,
         'Whole-body context guides a thoracic locator. A focused 3D classifier examines two high-resolution blocks to predict CoA-like change from attenuation and image validity.', 29)
    text('evaluation-heading', x, 8.78, cw, .55, 'GENERALIZATION BY ANATOMY', 31, ACCENTS[2], True)
    text('evaluation-copy', x, 9.48, cw, 2.05,
         'Hold out a related phantom pair and every sibling from that anatomy. Rotate across five development families; retain the infant pair for later testing.', 28)

    fig_titles = ('FIGURE A  |  BITWISE MORPHOLOGY',
                  'FIGURE B  |  VIRTUAL POPULATION',
                  'FIGURE C  |  ROC + FAMILY EVALUATION')
    fig_plans = (
        'Original → dilation / erosion\nBorrowed and released voxels\nNeighbor-based label reassignment',
        'Common-view 3D aortic variants\nBaseline → narrowing → companion changes\nParameter grid + achieved morphology',
        'Exact empirical ROC curves\nValidation and held-out results labeled separately\nFamily performance + control comparison',
    )
    captions = (
        'Figure A. Planned color-only pixel schematic: Boolean mask changes, a resistant neighbor and deterministic tissue-label reassignment.',
        'Figure B. Planned matched-view sibling examples. Use common orientation and scale; report achieved edits. Siblings are correlated, not independent patients.',
        'Figure C. Planned measured ROC comparison. Positive labels use ≥20% achieved local target-volume loss, an engineering proxy rather than clinical severity.',
    )
    for i, x in enumerate(cols):
        rectangle(f'figure-placeholder-{chr(65+i)}', x, 12.05, cw, 11.6, 'F3F6F8', 'A8BAC5', True)
        text(f'figure-title-{i}', x+.55, 13.05, cw-1.1, .75, fig_titles[i], 30, ACCENTS[i], True, 'center')
        text(f'figure-content-{i}', x+.6, 16.65, cw-1.2, 3.7, fig_plans[i], 31, GRAY, False, 'center')
        text(f'figure-status-{i}', x+.55, 22.1, cw-1.1, .65, 'ONE COMPOSITE FIGURE PLACEHOLDER', 22, GRAY, False, 'center')
        text(f'figure-caption-{i}', x, 24.05, cw, 1.6, captions[i], 24, GRAY)

    text('methods-heading', cols[0], 26.12, cw, .6, 'GEOMETRY + VOXEL OWNERSHIP', 31, ACCENTS[0], True)
    text('methods-copy', cols[0], 27, cw, 5.12,
         'Dilation expands a Boolean target mask into eligible neighboring voxels; protected labels resist borrowing.\n\n'
         'Erosion releases target voxels and assigns surrounding labels using weighted physical-neighborhood paths.\n\n'
         'Copy identity and attenuation from original seeds. Preserve unresolved releases; use optional connectivity and local-volume safeguards.', 28)
    text('cohort-method-heading', cols[1], 26.12, cw, .6, 'CONTROLLED COVERAGE + TRACEABILITY', 31, ACCENTS[1], True)
    text('cohort-method-copy', cols[1], 27, cw, 5.12,
         'Three dilation levels × five CoA erosion levels, plus an unmodified baseline per phantom.\n\n'
         'Freeze reviewed inputs, achieved changes, anatomical labels and paired attenuation overlays.\n\n'
         'Six provisional reference-anatomy families. Group related cases and siblings across training and evaluation.', 28)
    text('results-heading', cols[2], 26.12, cw, .6, 'PRELIMINARY EVIDENCE + NEXT STEP', 31, ACCENTS[2], True)
    text('results-copy', cols[2], 27, cw, 5.12,
         'Fixed-split validation: AUROC 0.875 (8 images; 50-epoch pilot).\n\n'
         'Family holdout at 200 epochs: pooled AUROC 0.531; balanced accuracy 50% (60 images). Generalization remains unresolved.\n\n'
         'Next: classification-only control, then matched population-size and diversity comparisons.', 28)

    rectangle('footer-rule', .95, 33.02, 46.1, .035, 'CEDBE3')
    text('references', .95, 33.35, 46.1, 1.04,
         '[1] Segars WP et al. 4D XCAT phantom for multimodality imaging research. Med Phys. 2010;37:4902–4915. doi:10.1118/1.3480985.\n'
         '[2] CDC. Coarctation of the Aorta. Updated Jan 8, 2026. https://www.cdc.gov/heart-defects/about/coarctation-of-the-aorta.html', 20, GRAY)
    text('scope-footer', .95, 35.05, 46.1, .46,
         'Engineering study • Synthetic attenuation proxies • Virtual-cohort benefit and clinical generalization remain research endpoints', 20, NAVY, True)
    return shapes


def wrapping(item):
    """Wrap the preview using the same physical text widths as the PowerPoint."""
    font = ImageFont.truetype(str(FONT_PATH), item['font'])
    limit = item['w'] * 72
    lines = []
    for paragraph in item['text'].split('\n'):
        if not paragraph:
            lines.append('')
            continue
        row = ''
        for word in paragraph.split():
            candidate = f'{row} {word}'.strip()
            if font.getlength(candidate) > limit and row:
                lines.append(row)
                row = word
            else:
                row = candidate
        lines.append(row)
    required = len(lines) * item['font'] / 72 * 1.10
    if required > item['h'] + .03:
        raise ValueError(f"Text overflows planned box: {item['name']} needs {required:.2f} in, has {item['h']}")
    return '\n'.join(lines)


def build(output):
    if output.exists():
        raise ValueError('Preserve existing output; choose a fresh revision')
    items = layout()
    assert sum(i['name'].startswith('figure-placeholder-') for i in items) == 3
    for i in items:
        assert 0 <= i['x'] and 0 <= i['y'] and i['x'] + i['w'] <= WIDTH + .01 and i['y'] + i['h'] <= HEIGHT
        if i['kind'] == 'text':
            i['wrapped_preview'] = wrapping(i)
    output.mkdir(parents=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(WIDTH), Inches(HEIGHT)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.core_properties.title = 'syncoCT BMES 2026 poster outline'
    presentation.core_properties.subject = 'Three-column 48 x 36 inch editable research poster'
    for i in items:
        if i['kind'] == 'rect':
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(i['x']), Inches(i['y']), Inches(i['w']), Inches(i['h']))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(i['fill'])
            if i['line']:
                shape.line.color.rgb = RGBColor.from_string(i['line']); shape.line.width = Pt(1.5)
                if i['dashed']: shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            else:
                shape.line.fill.background()
        else:
            shape = slide.shapes.add_textbox(Inches(i['x']), Inches(i['y']), Inches(i['w']), Inches(i['h']))
            frame = shape.text_frame; frame.clear()
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.TOP
            for n, line in enumerate(i['wrapped_preview'].split('\n')):
                paragraph = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                paragraph.text = line; paragraph.line_spacing = 1.10
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.alignment = PP_ALIGN.CENTER if i['align'] == 'center' else PP_ALIGN.LEFT
                paragraph.font.name = 'DejaVu Sans'; paragraph.font.size = Pt(i['font'])
                paragraph.font.bold = i['bold']; paragraph.font.color.rgb = RGBColor.from_string(i['color'])
        shape.name = i['name']
    slide.notes_slide.notes_text_frame.text = (
        'OUTLINE ONLY. Exactly one figure placeholder in each column.\n'
        'Read archive/research/docs/poster/bmes2026-poster-handoff.txt in the FakeCT fakect.26.09.16 worktree.\n'
        'The latest classification-only control was launched on 2026-10-05; no outcomes from it are claimed here.\n'
        '48 x 36 inches is the requested format, not a verified BMES 2026 size requirement.\n'
        'Image data are attenuation-copy proxies, not reconstructed clinical CT. Clinical representativeness is unvalidated.\n'
        'The fixed-split AUROC is development validation; grouped holdout uses five provisional anatomy families.\n'
        'Authors, affiliations and contact require user input. XCAT reference: https://scholars.duke.edu/publication/1255117\n'
        'CDC clinical background: https://www.cdc.gov/heart-defects/about/coarctation-of-the-aorta.html'
    )
    pptx = output/'fakect-bmes2026-outline.pptx'
    presentation.save(pptx)

    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    fig = plt.figure(figsize=(WIDTH, HEIGHT), dpi=50)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set(xlim=(0, WIDTH), ylim=(HEIGHT, 0)); ax.axis('off')
    for i in items:
        if i['kind'] == 'rect':
            ax.add_patch(Rectangle((i['x'], i['y']), i['w'], i['h'], facecolor='#'+i['fill'],
                                   edgecolor='#'+i['line'] if i['line'] else 'none',
                                   linewidth=1.5, linestyle='--' if i['dashed'] else '-'))
        else:
            ax.text(i['x']+i['w']/2 if i['align']=='center' else i['x'], i['y'], i['wrapped_preview'],
                    fontsize=i['font'], color='#'+i['color'], weight='bold' if i['bold'] else 'normal',
                    family='DejaVu Sans', va='top', ha=i['align'], linespacing=1.10)
    fig.savefig(output/'fakect-bmes2026-outline.png', dpi=50)
    fig.savefig(output/'fakect-bmes2026-outline.pdf', metadata={'CreationDate':None,'ModDate':None})
    plt.close(fig)
    (output/'poster-outline-layout.json').write_text(json.dumps({'width_inches':WIDTH,'height_inches':HEIGHT,'shapes':items}, indent=2))
    prompt = ROOT/'archive/research/docs/poster/bmes2026-poster-handoff.txt'
    (output/'poster-handoff.txt').write_bytes(prompt.read_bytes())
    sources = [Path(__file__), prompt,
               ROOT/'outputs/experiments/coa/focused-pilot-r1/pilot/results.json',
               ROOT/'outputs/experiments/coa/rotating-budget200-r1/review-r2/analysis.json',
               ROOT/'outputs/datasets/coa/context-aorta-r1/dataset-manifest.json']
    proof = {'status':'outline','slide_count':1,'width_inches':WIDTH,'height_inches':HEIGHT,
             'figure_placeholders':3,'editable_text_and_shapes':True,
             'preview_method':'Same physical layout and wrapped text as native PowerPoint; Matplotlib preview, not Office rendering',
             'new_training_results_claimed':False,'source_sha256':{str(p):digest(p) for p in sources},
             'web_sources':['https://scholars.duke.edu/publication/1255117','https://www.cdc.gov/heart-defects/about/coarctation-of-the-aorta.html'],
             'artifact_sha256':{p.name:digest(p) for p in output.iterdir() if p.is_file()}}
    (output/'outline-manifest.json').write_text(json.dumps(proof, indent=2))
    reopened = Presentation(pptx)
    assert len(reopened.slides) == 1
    assert reopened.slide_width == Inches(48) and reopened.slide_height == Inches(36)
    assert len([s for s in reopened.slides[0].shapes if s.name.startswith('figure-placeholder-')]) == 3
    assert all(s.shape_type != 13 for s in reopened.slides[0].shapes)  # Native outline, no flattened pictures.
    return proof


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='outputs/poster/bmes2026/outline-r1')
    arguments = parser.parse_args()
    path = Path(arguments.output)
    build(path if path.is_absolute() else ROOT/path)
    print(str((path if path.is_absolute() else ROOT/path)/'fakect-bmes2026-outline.pptx'))
